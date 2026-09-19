"""Reading submissions and writing annotated workbooks.

Output files are always new workbooks. The original submission is only ever
read.
"""

from __future__ import annotations

import io
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from railvalidator.review import Change, Finding, summarise

# Soft fills so the sheet stays readable when printed in greyscale.
VERDICT_FILLS = {
    "MATCH": None,
    "ALIAS": "DCE6F2",
    "TYPO": "FFF2CC",
    "AMBIGUOUS": "FCE4D6",
    "UNKNOWN": "F4CCCC",
}
_HEADER_FILL = PatternFill("solid", start_color="F2F2F2")
_HEADER_BORDER = Border(bottom=Side(style="thin", color="BFBFBF"))


def read_submission(source: Path | str | io.BytesIO) -> pd.DataFrame:
    return pd.read_excel(source, dtype=str).fillna("")


def _write_table(ws: Worksheet, df: pd.DataFrame) -> None:
    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.border = _HEADER_BORDER
    for row in df.itertuples(index=False):
        ws.append(list(row))
    for i, column in enumerate(df.columns, start=1):
        longest = max([len(str(column))] + [len(str(v)) for v in df[column]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 10), 70)
    ws.freeze_panes = "A2"


def _fill(verdict: str) -> PatternFill | None:
    colour = VERDICT_FILLS.get(verdict)
    return PatternFill("solid", start_color=colour) if colour else None


def findings_frame(findings: list[Finding]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "row": f.row + 2,  # spreadsheet row number, counting the header
                "column": f.column,
                "value": f.value,
                "verdict": f.verdict.value,
                "proposed": f.proposed_name or "",
                "candidates": "; ".join(c.name for c in f.result.candidates) if f.proposed_name is None else "",
                "explanation": f.explanation,
            }
            for f in findings
            if f.verdict.needs_review
        ],
        columns=["row", "column", "value", "verdict", "proposed", "candidates", "explanation"],
    )


def changes_frame(changes: list[Change]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"row": c.row + 2, "column": c.column, "old": c.old, "new": c.new, "verdict": c.verdict, "reason": c.reason}
            for c in changes
        ],
        columns=["row", "column", "old", "new", "verdict", "reason"],
    )


def build_report_workbook(df: pd.DataFrame, findings: list[Finding]) -> Workbook:
    """Validation sheet (colour-coded cells with comments), findings list and summary."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Validation"
    _write_table(ws, df)
    for f in findings:
        cell = ws.cell(row=f.row + 2, column=df.columns.get_loc(f.column) + 1)
        fill = _fill(f.verdict.value)
        if fill:
            cell.fill = fill
        if f.verdict.needs_review:
            cell.comment = Comment(f"{f.verdict.value}: {f.explanation}", "Rail Station Validator")

    _write_table(wb.create_sheet("Findings"), findings_frame(findings))

    summary = wb.create_sheet("Summary")
    counts = summarise(findings)
    _write_table(summary, pd.DataFrame({"verdict": list(counts), "cells": list(counts.values())}))
    for r in range(2, summary.max_row + 1):
        fill = _fill(summary.cell(row=r, column=1).value)
        if fill:
            summary.cell(row=r, column=1).fill = fill
    return wb


def build_commit_workbook(corrected: pd.DataFrame, changes: list[Change]) -> Workbook:
    """The corrected submission plus a log of every approved change."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Corrected"
    _write_table(ws, corrected)
    for c in changes:
        cell = ws.cell(row=c.row + 2, column=corrected.columns.get_loc(c.column) + 1)
        cell.fill = _fill(c.verdict) or PatternFill()
        cell.comment = Comment(f"Was '{c.old}'. {c.reason}", "Rail Station Validator")
    _write_table(wb.create_sheet("Change log"), changes_frame(changes))
    return wb


def workbook_bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
