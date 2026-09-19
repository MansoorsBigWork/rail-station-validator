import pandas as pd
from openpyxl import load_workbook

from railvalidator.excel_io import build_commit_workbook, build_report_workbook
from railvalidator.explain import explain
from railvalidator.review import apply_decisions, default_decisions, summarise, validate_submission


def submission():
    return pd.DataFrame(
        {
            "claim_id": ["EXP-1", "EXP-2", "EXP-3"],
            "from_station": ["Manchester Picadilly", "Manchester", "York"],
            "to_station": ["Lime St", "Leeds", "Atlantis"],
        }
    )


def test_summary_counts(catalog):
    counts = summarise(validate_submission(submission(), catalog))
    assert counts == {"MATCH": 2, "ALIAS": 1, "TYPO": 1, "AMBIGUOUS": 1, "UNKNOWN": 1}


def test_default_decisions_exclude_cells_needing_a_person(catalog):
    decisions = default_decisions(validate_submission(submission(), catalog))
    assert decisions == {(0, "from_station"): "Manchester Piccadilly", (0, "to_station"): "Liverpool Lime Street"}


def test_apply_never_modifies_the_original(catalog):
    df = submission()
    before = df.copy()
    findings = validate_submission(df, catalog)
    decisions = default_decisions(findings) | {(1, "from_station"): "Manchester Victoria"}
    corrected, changes = apply_decisions(df, findings, decisions)
    pd.testing.assert_frame_equal(df, before)
    assert corrected.loc[1, "from_station"] == "Manchester Victoria"
    assert [c.old for c in changes] == ["Manchester Picadilly", "Lime St", "Manchester"]


def test_every_explanation_names_its_evidence(catalog):
    for f in validate_submission(submission(), catalog):
        text = explain(f.result, catalog)
        assert f.value in text
        for cand in f.result.candidates[:1]:
            assert cand.name in text


def test_workbooks_are_written(catalog, tmp_path):
    df = submission()
    findings = validate_submission(df, catalog)
    build_report_workbook(df, findings).save(tmp_path / "report.xlsx")
    corrected, changes = apply_decisions(df, findings, default_decisions(findings))
    build_commit_workbook(corrected, changes).save(tmp_path / "commit.xlsx")

    assert load_workbook(tmp_path / "report.xlsx").sheetnames == ["Validation", "Findings", "Summary"]
    log = load_workbook(tmp_path / "commit.xlsx")["Change log"]
    assert log.max_row == 1 + len(changes)
